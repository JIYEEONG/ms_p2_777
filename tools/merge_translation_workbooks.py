"""Compare the two translation drafts and keep the selected copy in the main workbook.

Run from the repository root: python tools/merge_translation_workbooks.py
The source workbook with Sanskrit names is never changed.
"""

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "translations" / "moov-ko-en.xlsx"
SANSKRIT = ROOT / "translations" / "moov-ko-en_san.xlsx"

# A few labels read differently once their actual UI placement is considered.
OVERRIDES = {
    "AI 말동무": ("AI chat companion", "Distinct from the broader AI companion feature."),
    "감성": ("Atmospheric", "Adjective for the trip mood filter."),
    "공간 테마 예시": ("Space theme preview", "Singular preview image in the theme dialog."),
    "디즈니 플러스": ("Disney+", "Use the service's displayed brand name."),
    "상품 구매": ("Shop", "Short label for the primary tab."),
    "상품 찾기": ("Browse products", "Label for the product browsing tab."),
    "연결": ("Connected", "Connection status shown for a linked OTT account."),
    "연결 안 함": ("Not connected", "Connection status shown for an unlinked OTT account."),
    "적용 테마": ("Active theme", "Label above the currently active theme."),
}

POSES = {
    "앉은 산 자세", "앉은 고양이·소 자세", "앉은 측면 늘리기", "앉은 척추 비틀기",
    "전사 자세 2", "나무 자세", "의자 자세", "하이 런지", "리버스 워리어",
    "삼각 자세", "반달 자세", "무용수 자세", "팔각 자세",
}

# The connection-status label matched in both first drafts, so its pre-merge
# value was absent from the first version of the audit sheet.
ORIGINAL_BASELINE = {"연결": "Connect"}


def main():
    book = load_workbook(OUTPUT)
    sheet = book["UI and voice"]
    headers = {cell.value: cell.column - 1 for cell in sheet[1]}
    san = load_workbook(SANSKRIT, read_only=True, data_only=True)["UI and voice"]

    if sheet.max_row != san.max_row:
        raise ValueError("The workbooks have different numbers of entries")

    # Preserve the first draft if the merge is run again after the audit sheet exists.
    previous = {}
    if "비교 기록" in book:
        for row in book["비교 기록"].iter_rows(min_row=2, values_only=True):
            previous[row[0]] = row[2]
        del book["비교 기록"]

    audit = book.create_sheet("비교 기록")
    audit.append(["ID", "한국어", "기존 English", "산스크리트어 파일 English", "최종 English", "선택", "메모"])
    differences = 0
    audit_rows = 0
    overrides = 0
    poses = 0

    for row, candidate in zip(sheet.iter_rows(min_row=2), san.iter_rows(min_row=2, values_only=True)):
        identifier, korean, current = [cell.value for cell in row[:3]]
        location = row[headers["원본 위치"]].value
        if (identifier, korean, location) != (candidate[0], candidate[1], candidate[4]):
            raise ValueError(f"Row alignment mismatch at {identifier}")

        original = ORIGINAL_BASELINE.get(korean, previous.get(identifier, current))
        proposed = candidate[2]
        chosen, note = OVERRIDES.get(korean, (proposed, ""))
        if korean in POSES:
            if "(" not in chosen or ")" not in chosen:
                raise ValueError(f"Sanskrit name missing from {korean}")
            note = "Yoga pose name includes Sanskrit."
            poses += 1

        row[2].value = chosen
        row[headers["검수 상태"]].value = candidate[3]
        if chosen != proposed:
            overrides += 1
        if original != proposed:
            differences += 1
        if original != proposed or chosen != proposed:
            selection = "기존" if chosen == original else "산스크리트어 파일" if chosen == proposed else "문맥 보정"
            audit.append([identifier, korean, original, proposed, chosen, selection, note])
            audit_rows += 1

    audit.freeze_panes = "C2"
    audit.auto_filter.ref = audit.dimensions
    for column, width in {"A": 16, "B": 55, "C": 58, "D": 62, "E": 62, "F": 20, "G": 58}.items():
        audit.column_dimensions[column].width = width
    for cell in audit[1]:
        cell.fill = PatternFill("solid", fgColor="234336")
        cell.font = Font(color="FFFFFF", bold=True)
    for row in audit.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    book.save(OUTPUT)
    print(f"Compared {sheet.max_row - 1} entries: {differences} differences, {audit_rows} audit rows, {overrides} contextual overrides, {poses} Sanskrit pose names")


if __name__ == "__main__":
    main()
