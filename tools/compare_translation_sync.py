"""Compare current translations with committed main and the previous English branch."""
import argparse
from io import BytesIO
import json
from pathlib import Path
import subprocess

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def git(*args):
    return subprocess.check_output(["git", *args], cwd=ROOT)


def sheets(workbook):
    return {name: {row[1]: row[2] or "" for row in list(workbook[name].values)[1:] if row[1]}
            for name in ("UI and voice", "Home rental") if name in workbook}


def comparison(old, current):
    return {
        "previous_count": len(old), "current_count": len(current),
        "untranslated": [key for key, value in current.items() if not value],
        "added": sorted(current.keys() - old.keys()),
        "removed": sorted(old.keys() - current.keys()),
        "changed_english": {key: {"before": old[key], "after": current[key]}
                            for key in sorted(old.keys() & current.keys()) if old[key] != current[key]},
    }


def dictionary(source):
    return json.loads(source.split("Object.freeze(", 1)[1].rsplit(");", 1)[0])


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base", default="origin/main")
    parser.add_argument("--english", default="english")
    args = parser.parse_args()
    old = sheets(load_workbook(BytesIO(git("show", f"{args.base}:translations/moov-ko-en.xlsx")), read_only=True))
    current = sheets(load_workbook(ROOT / "translations/moov-ko-en.xlsx", read_only=True))
    previous_dictionary = dictionary(git("show", f"{args.english}:front/public/i18n-data.js").decode("utf-8"))
    current_dictionary = dictionary((ROOT / "front/public/i18n-data.js").read_text(encoding="utf-8"))
    report = {
        "base_commit": git("rev-parse", args.base).decode().strip(),
        "previous_english_commit": git("rev-parse", args.english).decode().strip(),
        "inventory_note": "Counts include extraction improvements: complete templates replace partial fragments; icon interpolations are omitted.",
        "workbook_vs_base": {name: comparison(old.get(name, {}), values) for name, values in current.items()},
        "shared_dictionary_vs_previous_english": comparison(previous_dictionary, current_dictionary),
    }
    output = ROOT / "translations/sync-report.json"
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    for name, data in report["workbook_vs_base"].items():
        print(f"{name}: {data['current_count']} current, {len(data['added'])} added, {len(data['removed'])} removed, {len(data['changed_english'])} revised, {len(data['untranslated'])} blank")
    data = report["shared_dictionary_vs_previous_english"]
    print(f"Previous English dictionary: {len(data['added'])} added, {len(data['removed'])} removed, {len(data['changed_english'])} revised")


if __name__ == "__main__":
    main()
