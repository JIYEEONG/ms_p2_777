"""Build a reviewable inventory of Korean UI and spoken text.

Run from the repository root: python tools/extract_translations.py
The workbook is deliberately separate from the application source.
"""

from collections import defaultdict
from html import unescape
from html.parser import HTMLParser
from pathlib import Path
import json
import re
import subprocess

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
FRONT = ROOT / "front"
OUTPUT = ROOT / "translations" / "moov-ko-en.xlsx"
HANGUL = re.compile(r"[가-힣]")
HTML_TEXT = re.compile(r">([^<>]*[가-힣][^<>]*)<")
HTML_ATTRIBUTE = re.compile(r'(?:alt|title|placeholder|aria-label)=["\']([^"\']+)["\']')
TAG = re.compile(r"<[^>]+>")
INTERPOLATION = re.compile(r"\$\{[^{}]*\}")
WHITESPACE = re.compile(r"\s+")


def fragments(raw):
    raw = unescape(raw.replace("\\n", " "))
    candidates = (HTML_TEXT.findall(raw) + HTML_ATTRIBUTE.findall(raw)) if "<" in raw else [raw]
    for candidate in candidates:
        candidate = TAG.sub(" ", candidate)
        candidate = INTERPOLATION.sub("{value}", candidate)
        candidate = WHITESPACE.sub(" ", candidate).strip().strip('"\'<> ')
        if candidate.startswith(("./", "/", "assets/")):
            continue
        if HANGUL.search(candidate) and len(candidate) <= 350:
            yield candidate


class VisibleHTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self.hidden = 0
        self.entries = []

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style"}:
            self.hidden += 1
        if self.hidden:
            return
        for name, value in attrs:
            if name in {"alt", "title", "placeholder", "aria-label"} and value:
                self.entries.append((value, self.getpos()[0]))

    def handle_endtag(self, tag):
        if tag in {"script", "style"}:
            self.hidden = max(0, self.hidden - 1)

    def handle_data(self, data):
        if not self.hidden and HANGUL.search(data):
            self.entries.append((data, self.getpos()[0]))


def inventory():
    entries = defaultdict(set)
    paths = [FRONT / "app.js", FRONT / "index.html", FRONT / "public" / "moov.html"]
    paths += sorted((FRONT / "public" / "space-content").rglob("*.js"))
    paths += sorted((FRONT / "public" / "space-content").rglob("*.html"))
    for path in paths:
        source = path.read_text(encoding="utf-8")
        relative = path.relative_to(ROOT).as_posix()
        js_entries = json.loads(subprocess.check_output(
            ["node", str(ROOT / "tools" / "extract_js_strings.cjs"), str(path)], cwd=ROOT
        ))
        for entry in js_entries:
            line = source.count("\n", 0, entry["position"]) + 1
            for phrase in fragments(entry["text"]):
                entries[phrase].add(f"{relative}:{line}")
        if path.suffix == ".html":
            parser = VisibleHTML()
            parser.feed(source)
            for value, line in parser.entries:
                for phrase in fragments(value):
                    entries[phrase].add(f"{relative}:{line}")
    return entries


def main():
    entries = inventory()
    translations = json.loads((ROOT / "translations" / "wellness-en.json").read_text(encoding="utf-8"))
    existing = {}
    if OUTPUT.exists():
        old_sheet = load_workbook(OUTPUT, read_only=True).active
        existing = {row[1]: (row[2], row[3]) for row in list(old_sheet.values)[1:] if row[1] and row[2]}
    OUTPUT.parent.mkdir(exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "UI and voice"
    sheet.append(["ID", "한국어", "English", "검수 상태", "원본 위치"])
    for index, (korean, locations) in enumerate(sorted(entries.items()), start=1):
        english, status = existing.get(korean, (translations.get(korean, ""), "검수 필요" if korean in translations else "번역 필요"))
        sheet.append([f"MOOV-{index:04d}", korean, english, status, "\n".join(sorted(locations))])
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {"A": 16, "B": 66, "C": 66, "D": 16, "E": 64}.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="234336")
        cell.font = Font(color="FFFFFF", bold=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(OUTPUT)
    translated = sum(bool(existing.get(phrase, (translations.get(phrase),))[0]) for phrase in entries)
    print(f"{len(entries)} Korean entries, {translated} translated -> {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
