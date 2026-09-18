"""Inventory text in the integrated home/rental module.

Run from the repository root: python tools/extract_home_translations.py
Existing English edits in the "Home rental" sheet are retained.
"""

from collections import defaultdict
from pathlib import Path
import json
import re
import subprocess

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from extract_translations import ROOT, VisibleHTML, fragments


MODULE = ROOT / "front" / "public" / "moov-home"
WORKBOOK = ROOT / "translations" / "moov-ko-en.xlsx"
SHEET_NAME = "Home rental"
SOURCE_FILES = [
    MODULE / "index.html",
    *(MODULE / "js" / name for name in (
        "data.js", "home.js", "rental-policy.js", "road-routing.js",
        "secure-cleanup.js", "shared.js", "state.js",
    )),
    *sorted((MODULE / "assets").glob("*.svg")),
]


def inventory():
    entries = defaultdict(set)
    for path in SOURCE_FILES:
        source = path.read_text(encoding="utf-8")
        relative = path.relative_to(ROOT).as_posix()
        if path.suffix == ".js":
            parsed = json.loads(subprocess.check_output(
                ["node", str(ROOT / "tools" / "extract_js_strings.cjs"), str(path)], cwd=ROOT
            ))
            for entry in parsed:
                line = source.count("\n", 0, entry["position"]) + 1
                for phrase in fragments(entry["text"]):
                    entries[phrase].add(f"{relative}:{line}")
        else:
            parser = VisibleHTML()
            parser.feed(source)
            for value, line in parser.entries:
                for phrase in fragments(value):
                    entries[phrase].add(f"{relative}:{line}")

    vehicle_file = MODULE / "js" / "data.js"
    vehicle_source = vehicle_file.read_text(encoding="utf-8")
    for match in re.finditer(r'name: "MOOV ([^\"]+)"', vehicle_source):
        entries[match.group(1)].add(
            f"{vehicle_file.relative_to(ROOT).as_posix()}:{vehicle_source.count(chr(10), 0, match.start()) + 1}"
        )

    dynamic_phrases = {
        "+{amount}원": "js/home.js",
        "✓ {vehicle}": "js/rental-policy.js",
        "{hours}시간 · 차량 요금": "js/rental-policy.js",
        "시간당 약 {amount}원": "js/rental-policy.js",
        "차량 {vehicleFare}원 + 옵션 {optionFare}원": "js/home.js",
        "{place} 주변의 {vehicle} 차량을 확인하고 있어요.": "js/rental-policy.js",
        "약 3초 후 코스 지도로 이동합니다.": "js/rental-policy.js",
        "{place} 방면 · 다음 지점 약 {minutes}분": "js/home.js",
        "{distance} km · 약 {minutes}분": "js/home.js",
        "· 약 {minutes}분": "js/home.js",
        "장소 목록": "js/home.js",
        "선택 후보: {place}": "js/rental-policy.js",
        "{number}. {place}": "js/rental-policy.js",
    }
    for phrase, source in dynamic_phrases.items():
        entries[phrase].add(f"{(MODULE / source).relative_to(ROOT).as_posix()} (dynamic)")

    # The large OSM JSON is source data, not app copy. Only inventory labels
    # in the small UI code following it; the geographic names remain in OSM.
    path = MODULE / "js" / "local-map.js"
    source = path.read_text(encoding="utf-8")
    tail = source.split("const MOOV_MAP_BOUNDS", 1)[1]
    for match in re.finditer(r"['\"]([^'\"\n]*[가-힣][^'\"\n]*)['\"]", tail):
        for phrase in fragments(match.group(1)):
            entries[phrase].add(f"{path.relative_to(ROOT).as_posix()}:{source.count(chr(10), 0, source.index('const MOOV_MAP_BOUNDS')) + tail.count(chr(10), 0, match.start()) + 1}")
    return entries


def main():
    entries = inventory()
    workbook = load_workbook(WORKBOOK)
    shared = {row[1]: row[2] for row in list(workbook["UI and voice"].values)[1:] if row[1] and row[2]}
    existing = {}
    if SHEET_NAME in workbook:
        existing = {row[1]: (row[2], row[3]) for row in list(workbook[SHEET_NAME].values)[1:] if row[1]}
        del workbook[SHEET_NAME]
    sheet = workbook.create_sheet(SHEET_NAME)
    sheet.append(["ID", "한국어", "English", "검수 상태", "원본 위치"])
    for index, (korean, locations) in enumerate(sorted(entries.items()), start=1):
        english, status = existing.get(korean, (shared.get(korean, ""), "기존 번역" if korean in shared else "번역 필요"))
        sheet.append([f"HOME-{index:04d}", korean, english, status, "\n".join(sorted(locations))])
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {"A": 16, "B": 70, "C": 70, "D": 18, "E": 72}.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="234336")
        cell.font = Font(color="FFFFFF", bold=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(WORKBOOK)
    print(f"Home rental: {len(entries)} entries, {sum(bool(sheet.cell(i, 3).value) for i in range(2, sheet.max_row + 1))} English translations")


if __name__ == "__main__":
    main()
