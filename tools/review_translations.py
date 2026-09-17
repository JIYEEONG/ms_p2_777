"""Apply MOOV terminology and flag translation rows that need attention."""

import json
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parents[1]
path = ROOT / "translations" / "moov-ko-en.xlsx"
corrections = json.loads((ROOT / "translations" / "ui-corrections.json").read_text(encoding="utf-8"))
book = load_workbook(path)
sheet = book.active
changed = 0
flagged = 0

for row in sheet.iter_rows(min_row=2):
    korean = row[1].value or ""
    english = row[2].value or ""
    replacement = corrections.get(korean)
    if replacement is None:
        minutes = re.fullmatch(r"(\d+)분", korean)
        remaining = re.fullmatch(r"(\d+)분 남음", korean)
        people = re.fullmatch(r"(\d+)~(\d+)인", korean)
        refrigerator = re.fullmatch(r"냉장함 ([A-Z]-\d+)", korean)
        locker = re.fullmatch(r"수납함 ([A-Z]-\d+)", korean)
        if minutes:
            replacement = f"{minutes[1]} min"
        elif remaining:
            replacement = f"{remaining[1]} min remaining"
        elif people:
            replacement = f"{people[1]}–{people[2]} people"
        elif refrigerator:
            replacement = f"Chilled Compartment {refrigerator[1]}"
        elif locker:
            replacement = f"Compartment {locker[1]}"
        elif korean.startswith("MOOV "):
            replacement = re.sub(r"\bMOV\b", "MOOV", english)
    if replacement and replacement != english:
        row[2].value = replacement
        row[3].value = "용어 초안 · 검수 필요"
        english = replacement
        changed += 1
    if re.search(r"[가-힣]", english) or not english.strip() or len(english) > max(100, 4 * len(korean)):
        row[3].value = "우선 검수 필요"
        for cell in row:
            cell.fill = PatternFill("solid", fgColor="FFF1D6")
        flagged += 1

book.save(path)
print(f"Corrected {changed} terms; flagged {flagged} rows")
