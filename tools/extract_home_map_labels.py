"""Inventory visible OSM map labels in the home translation workbook.

Run from the repository root: python tools/extract_home_map_labels.py
Romanized names are drafts for editorial review; source OSM data is unchanged.
"""

from collections import defaultdict
import json
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
MAP = ROOT / "front" / "public" / "moov-home" / "js" / "local-map.js"
WORKBOOK = ROOT / "translations" / "moov-ko-en.xlsx"
SHEET = "Map labels"
INITIAL = ["g", "kk", "n", "d", "tt", "r", "m", "b", "pp", "s", "ss", "", "j", "jj", "ch", "k", "t", "p", "h"]
MEDIAL = ["a", "ae", "ya", "yae", "eo", "e", "yeo", "ye", "o", "wa", "wae", "oe", "yo", "u", "wo", "we", "wi", "yu", "eu", "ui", "i"]
FINAL = ["", "k", "k", "k", "n", "n", "n", "t", "l", "k", "m", "l", "l", "l", "p", "l", "m", "p", "p", "t", "t", "ng", "t", "t", "k", "t", "p", "t"]
CURATED = {
    "한강": "Han River", "청계천": "Cheonggyecheon Stream", "중랑천": "Jungnang Stream",
    "서울숲": "Seoul Forest", "서울어린이대공원": "Seoul Children's Grand Park",
    "반포한강공원": "Banpo Hangang Park", "뚝섬한강공원": "Ttukseom Hangang Park",
    "이촌한강공원": "Ichon Hangang Park", "잠원한강공원": "Jamwon Hangang Park",
    "서울광장": "Seoul Plaza", "낙산공원": "Naksan Park", "도산공원": "Dosan Park",
}


def romanize(text):
    result = []
    for character in text:
        code = ord(character) - 0xAC00
        if 0 <= code <= 11171:
            result.append(INITIAL[code // 588] + MEDIAL[(code % 588) // 28] + FINAL[code % 28])
        else:
            result.append(character)
    return re.sub(r"\b[a-z]", lambda match: match.group().upper(), "".join(result))


def draft_name(name):
    if name in CURATED:
        return CURATED[name], "번역 초안"
    suffixes = (
        ("어린이공원", " Children's Park"),
        ("근린공원", " Neighborhood Park"),
        ("한강공원", " Hangang Park"),
        ("공원", " Park"),
        ("광장", " Plaza"),
        ("대교", " Bridge"),
        ("연못", " Pond"),
        ("호수", " Lake"),
        ("숲", " Forest"),
        ("천", " Stream"),
    )
    for korean, english in suffixes:
        if name.endswith(korean) and len(name) > len(korean):
            return romanize(name[:-len(korean)]).strip() + english, "로마자 표기·감수 필요"
    return romanize(name), "로마자 표기·감수 필요"


def main():
    source = MAP.read_text(encoding="utf-8")
    data = json.loads(source[source.index("{"):source.index(";\nconst MOOV_MAP_BOUNDS")])
    labels = defaultdict(set)
    for feature in data["features"]:
        properties = feature["properties"]
        name = properties.get("name") or ""
        if not properties.get("label") or not re.search("[가-힣]", name):
            continue
        if properties["kind"] == "rail" or (properties["kind"] == "road" and properties["class"] not in {"primary", "secondary", "tertiary"}):
            continue
        labels[name].add((properties["osm_id"], properties["kind"]))

    workbook = load_workbook(WORKBOOK)
    shared = {row[1]: row[2] for sheet_name in ("UI and voice", "Home rental")
              for row in list(workbook[sheet_name].values)[1:] if row[1] and row[2]}
    previous = {}
    if SHEET in workbook:
        previous = {row[1]: (row[2], row[3]) for row in list(workbook[SHEET].values)[1:] if row[1]}
        del workbook[SHEET]
    sheet = workbook.create_sheet(SHEET)
    sheet.append(["ID", "한국어", "English", "검수 상태", "OSM ID", "종류"])
    for index, (name, references) in enumerate(sorted(labels.items()), start=1):
        if name in previous:
            english, status = previous[name]
        elif name in shared:
            english, status = shared[name], "기존 번역"
        else:
            english, status = draft_name(name)
        osm, kind = sorted(references)[0]
        sheet.append([f"MAP-{index:04d}", name, english, status, osm, kind])
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {"A": 15, "B": 42, "C": 48, "D": 24, "E": 25, "F": 12}.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="234336")
        cell.font = Font(color="FFFFFF", bold=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(WORKBOOK)
    print(f"Map labels: {len(labels)} visible Korean names")


if __name__ == "__main__":
    main()
