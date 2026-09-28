from pathlib import Path
from docx import Document

doc = Document('.tmp/storyboard/MOOV_관리자_페이지_기획서_최종.docx')
lines = []
for p in doc.paragraphs:
    t = p.text.strip()
    if t:
        lines.append(t)
for i, table in enumerate(doc.tables, 1):
    lines.append(f'\n[TABLE {i}]')
    for row in table.rows:
        lines.append(' || '.join(cell.text.replace('\n', ' / ').strip() for cell in row.cells))
Path('.tmp/storyboard/docx_text.txt').write_text('\n'.join(lines), encoding='utf-8')
